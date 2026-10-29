"""가상 증례 30건으로 프로그램을 검증한다.

  python vignette_validation.py template   # 평가자용 엑셀(validation/평가지.xlsx) 만들기
  python vignette_validation.py program    # 프로그램 판정 저장(validation/program_output.csv) — 평가가 끝난 뒤에 보기
  python vignette_validation.py score validation/평가지.xlsx   # 평가자 간 일치도, 프로그램 vs 합의 정답

평가 방식: 평가자는 증례 글만 읽고 '필요한 주의 범주'에 O 표시. 두 평가자의 판단이 다르면 토의(또는 교수님 확인)로
'합의정답' 시트를 채운다. 프로그램 판정은 규칙 ID를 같은 범주로 묶어 비교한다.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import yaml

from engine import assess

ROOT = Path(__file__).parent
VIG = ROOT / "validation" / "vignettes.yaml"

# 평가 범주 → 해당하는 규칙 ID
CATEGORIES: dict[str, list[str]] = {
    "실신 위험": ["A1", "A2", "A3", "A4", "A7"],
    "출혈 위험": ["A5", "A6", "Q3"],
    "저혈당 위험": ["B1", "B2", "B3"],
    "임신 확인 필요": ["E1", "E2"],
    "낙상 위험": ["F1", "F2"],
    "의사소통 지원": ["I3", "G5"],
    "공복 미충족": ["Q0", "Q1"],
    "생리 중 소변검사": ["Q2"],
    "발열·감염": ["I2"],
    "혈압 매우 높음": ["P1", "G3", "H2"],
}
CATS = list(CATEGORIES)
SHEETS = ["평가자A", "평가자B", "합의정답"]


def load_vignettes() -> tuple[list[dict], list[str]]:
    with open(VIG, encoding="utf-8") as f:
        v = yaml.safe_load(f)
    return v["cases"], v["exams"]


def program_categories(fields: dict, exams: list[str]) -> tuple[set[str], list[str]]:
    ids = [a.id for a in assess(fields, exams)]
    cats = {c for c, rules in CATEGORIES.items() if set(rules) & set(ids)}
    return cats, ids


def program_table() -> pd.DataFrame:
    cases, exams = load_vignettes()
    rows = []
    for c in cases:
        cats, ids = program_categories(c["fields"], exams)
        rows.append({"id": c["id"], **{k: k in cats for k in CATS}, "규칙": ", ".join(ids)})
    return pd.DataFrame(rows).set_index("id")


# ---------------------------------------------------------------- 엑셀 평가지
def make_template(path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    cases, _ = load_vignettes()
    wb = Workbook()
    font = "맑은 고딕"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E46")
    input_fill = PatternFill("solid", fgColor="FFF2CC")

    ws = wb.active
    ws.title = "안내"
    guide = [
        "검진 안전 알리미 — 증례 평가지",
        "",
        "1. '평가자A', '평가자B' 시트를 서로 다른 사람이 각자 채웁니다. 서로의 답이나 프로그램 결과를 보지 마세요.",
        "2. 증례 글만 읽고, 검진 중 직원이 특별히 주의하거나 조치해야 할 범주에 O를 입력합니다(노란 칸). 해당 없으면 비워 둡니다.",
        "3. 두 평가자의 답이 다른 칸은 토의하거나 지도교수님께 확인해 '합의정답' 시트에 최종 답을 적습니다.",
        "4. 다 채운 뒤: python vignette_validation.py score validation/평가지.xlsx",
        "",
        "범주 설명",
        "실신 위험: 채혈 중·후 혈관미주신경성 실신 가능성이 높아 자세·관찰 조치가 필요",
        "출혈 위험: 채혈·구강검진 후 지혈에 신경 써야 함",
        "저혈당 위험: 금식 중 저혈당 가능성이 있어 혈당 확인·채혈 순서 조정이 필요",
        "임신 확인 필요: 흉부 X-ray 전에 임신 여부 확인·보류가 필요",
        "낙상 위험: 이동·자세 변경 때 보조가 필요",
        "의사소통 지원: 인지·청력 문제로 지시 전달에 도움이 필요",
        "공복 미충족: 공복혈당·지질 결과 해석에 문제가 생길 정도로 공복이 짧음",
        "생리 중 소변검사: 소변검사 결과가 왜곡될 수 있음",
        "발열·감염: 급성 증상으로 검진 진행 여부를 의료진이 판단해야 함",
        "혈압 매우 높음: 재측정·의사 보고가 필요한 수준",
        "",
        "예시) V00 증례에서 실신 위험만 해당하면 '실신 위험' 칸에만 O, 나머지는 빈칸.",
    ]
    for i, t in enumerate(guide, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=font, size=14 if i == 1 else 11, bold=i in (1, 8))
    ws.column_dimensions["A"].width = 120

    dv_proto = '"O"'
    for name in SHEETS:
        sh = wb.create_sheet(name)
        headers = ["증례", "증례 내용"] + CATS
        for j, h in enumerate(headers, 1):
            c = sh.cell(row=1, column=j, value=h)
            c.font = Font(name=font, bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        dv = DataValidation(type="list", formula1=dv_proto, allow_blank=True,
                            error="O 또는 빈칸만 입력하세요.", errorTitle="입력 확인")
        sh.add_data_validation(dv)
        # 예시 행
        ex = ["V00(예시)", "예시: 24세 여성, 과거 채혈 중 실신, 혈압 90/58 → 실신 위험에만 O"] + ["O" if k == "실신 위험" else "" for k in CATS]
        for j, v in enumerate(ex, 1):
            c = sh.cell(row=2, column=j, value=v)
            c.font = Font(name=font, italic=True, color="7F7F7F")
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if j > 2 else "left")
            c.border = border
        for i, case in enumerate(cases, 3):
            sh.cell(row=i, column=1, value=case["id"]).font = Font(name=font, bold=True)
            n = sh.cell(row=i, column=2, value=case["narrative"])
            n.font = Font(name=font)
            n.alignment = Alignment(wrap_text=True, vertical="top")
            for j in range(3, 3 + len(CATS)):
                c = sh.cell(row=i, column=j)
                c.fill = input_fill
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = Font(name=font, bold=True)
                dv.add(c)
            for j in range(1, 3 + len(CATS)):
                sh.cell(row=i, column=j).border = border
            sh.row_dimensions[i].height = 48
        last = 2 + len(cases)
        total_row = last + 2
        sh.cell(row=total_row, column=2, value="범주별 O 개수").font = Font(name=font, bold=True)
        for j in range(3, 3 + len(CATS)):
            col = sh.cell(row=1, column=j).column_letter
            c = sh.cell(row=total_row, column=j, value=f'=COUNTIF({col}3:{col}{last},"O")')
            c.font = Font(name=font, bold=True)
            c.alignment = Alignment(horizontal="center")
        sh.column_dimensions["A"].width = 10
        sh.column_dimensions["B"].width = 70
        for j in range(3, 3 + len(CATS)):
            sh.column_dimensions[sh.cell(row=1, column=j).column_letter].width = 11
        sh.row_dimensions[1].height = 36
        sh.freeze_panes = "C3"
    path.parent.mkdir(exist_ok=True)
    wb.save(path)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame | None:
    try:
        df = pd.read_excel(path, sheet_name=sheet, dtype=str)
    except ValueError:
        return None
    df = df[df["증례"].astype(str).str.match(r"^V\d\d$")].set_index("증례")
    if df[CATS].isna().all().all():
        return None  # 아직 비어 있음
    return df[CATS].apply(lambda s: s.fillna("").str.strip().str.upper().isin(["O", "1", "○", "V", "Y"]))


# ---------------------------------------------------------------- 지표
def kappa(a: pd.Series, b: pd.Series) -> float:
    a, b = a.astype(bool), b.astype(bool)
    po = (a == b).mean()
    pe = a.mean() * b.mean() + (1 - a.mean()) * (1 - b.mean())
    return float("nan") if pe == 1 else (po - pe) / (1 - pe)


def compare(truth: pd.DataFrame, pred: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for c in CATS:
        t, p = truth[c], pred.loc[truth.index, c]
        tp, fp = int((t & p).sum()), int((~t & p).sum())
        fn, tn = int((t & ~p).sum()), int((~t & ~p).sum())
        rows.append({"범주": c, "정답 O": int(t.sum()), "TP": tp, "FP": fp, "FN": fn, "TN": tn,
                     "민감도": tp / (tp + fn) if tp + fn else float("nan"),
                     "특이도": tn / (tn + fp) if tn + fp else float("nan"),
                     "양성예측도": tp / (tp + fp) if tp + fp else float("nan")})
    df = pd.DataFrame(rows)
    T, Pp = truth[CATS].stack(), pred.loc[truth.index, CATS].stack()
    tp, fp = int((T & Pp).sum()), int((~T & Pp).sum())
    fn, tn = int((T & ~Pp).sum()), int((~T & ~Pp).sum())
    df.loc[len(df)] = {"범주": "전체", "정답 O": int(T.sum()), "TP": tp, "FP": fp, "FN": fn, "TN": tn,
                       "민감도": tp / (tp + fn) if tp + fn else float("nan"),
                       "특이도": tn / (tn + fp) if tn + fp else float("nan"),
                       "양성예측도": tp / (tp + fp) if tp + fp else float("nan")}
    return df


def score(path: Path) -> dict:
    prog = program_table()
    out: dict = {"program": prog}
    ra, rb, cons = (read_sheet(path, s) for s in SHEETS)
    if ra is not None and rb is not None:
        out["rater_kappa"] = kappa(ra[CATS].stack(), rb[CATS].stack())
        out["rater_agreement"] = float((ra[CATS] == rb[CATS]).stack().mean())
    truth = cons if cons is not None else None
    if truth is not None:
        out["vs_truth"] = compare(truth, prog)
        out["prog_kappa"] = kappa(truth[CATS].stack(), prog.loc[truth.index, CATS].stack())
        disagree = []
        for cid in truth.index:
            for c in CATS:
                if truth.at[cid, c] != prog.at[cid, c]:
                    disagree.append({"증례": cid, "범주": c, "정답": "O" if truth.at[cid, c] else "",
                                     "프로그램": "O" if prog.at[cid, c] else "", "프로그램 규칙": prog.at[cid, "규칙"]})
        out["disagreements"] = pd.DataFrame(disagree)
    return out


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "program"
    if cmd == "template":
        p = ROOT / "validation" / "평가지.xlsx"
        make_template(p)
        print(f"평가지 생성: {p}")
    elif cmd == "program":
        t = program_table()
        t.replace({True: "O", False: ""}).to_csv(ROOT / "validation" / "program_output.csv", encoding="utf-8-sig")
        print(t.replace({True: "O", False: ""}).to_string())
    elif cmd == "score":
        r = score(Path(sys.argv[2]))
        if "rater_kappa" in r:
            print(f"평가자 간 일치율 {r['rater_agreement']:.1%}, Cohen's kappa {r['rater_kappa']:.2f}")
        if "vs_truth" in r:
            print(r["vs_truth"].round(2).to_string(index=False))
            print(f"프로그램 vs 합의정답 Cohen's kappa {r['prog_kappa']:.2f}")
            print("\n불일치 목록\n" + r["disagreements"].to_string(index=False))
        else:
            print("'합의정답' 시트가 비어 있습니다.")
