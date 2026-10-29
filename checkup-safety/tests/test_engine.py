"""규칙 엔진 회귀 테스트: pytest -q"""
import sys
from pathlib import Path

import pytest
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from engine import CORE_EXAMS, EXAMS, STATIONS, assess, derive, recommend_order, rulebook  # noqa: E402

CASES = yaml.safe_load(open(ROOT / "demo_cases.yaml", encoding="utf-8"))["cases"]


def ids(fields, exams):
    return {a.id for a in assess(fields, exams)}


# ---------- 1. 규칙 파일 자체 점검 ----------
def test_rulebook_integrity():
    book = rulebook()
    rule_ids = [r["id"] for r in book["rules"]]
    assert len(rule_ids) == len(set(rule_ids)), "규칙 ID 중복"
    for r in book["rules"]:
        assert r["level"] in ("red", "yellow")
        assert r["actions"], f"{r['id']} 조치 없음"
        assert r.get("evidence"), f"{r['id']} 근거 없음"
        for e in r["evidence"]:
            assert e in book["references"], f"{r['id']}: 참고문헌 {e} 없음"
        for s in r.get("suppressed_by", []):
            assert s in rule_ids


# ---------- 2. 시나리오 테스트 ----------
@pytest.mark.parametrize("case", CASES, ids=[c["name"] for c in CASES])
def test_demo_cases(case):
    fired = ids(case["fields"], case["exams"])
    for e in case.get("expect", []):
        assert e in fired, f"{e} 가 떠야 함. 실제: {sorted(fired)}"
    for a in case.get("absent", []):
        assert a not in fired, f"{a} 가 뜨면 안 됨. 실제: {sorted(fired)}"
    if "order_second" in case:
        order, _ = recommend_order(case["exams"], assess(case["fields"], case["exams"]))
        assert order[1] == case["order_second"], order


# ---------- 3. 경계값 테스트 ----------
@pytest.mark.parametrize("glucose,expected", [(69, True), (70, False), (54, True)])
def test_hypoglycemia_threshold(glucose, expected):
    assert ("B2" in ids({"glucose_poc": glucose}, [])) == expected


@pytest.mark.parametrize("egfr,expected", [(29.9, True), (30, False), (45, False)])
def test_egfr_threshold(egfr, expected):
    assert ("C1" in ids({"egfr": egfr, "age": 50}, ["ct_contrast"])) == expected


def test_bmi_calculation():
    p = derive({"height_cm": 160, "weight_kg": 46}, [])
    assert p["bmi"] == 18.0


def test_bmi_boundary_underweight():
    # BMI 18.5 정확히 = 저체중 아님
    f = {"height_cm": 200, "weight_kg": 74, "sbp": 95}
    assert derive(f, [])["bmi"] == 18.5
    assert "A3" not in ids(f, ["blood_draw"])


def test_stopbang_partial():
    p = derive({"snoring": True, "age": 60, "sex": "M"}, [])
    assert p["stopbang"] == 3 and p["stopbang_items"] == 3


# ---------- 4. 검사 의존성: 검사가 없으면 경고도 없어야 한다 ----------
def test_exam_required():
    f = {"implant": True, "egfr": 20, "prior_syncope_blood": True, "sbp": 85}
    fired = ids(f, [])
    assert not fired & {"D1", "C1", "A1"}


def test_suppression():
    f = {"prior_syncope_blood": True, "sbp": 85}
    fired = ids(f, ["blood_draw"])
    assert "A1" in fired and "A2" not in fired and "A4" not in fired


# ---------- 5. 결측 처리: 정보가 없을 때 '위험 없음'으로 단정하지 않는다 ----------
def test_missing_egfr_prompts_check():
    assert "C2" in ids({"age": 70}, ["ct_contrast"])


def test_missing_pregnancy_prompts_check():
    assert "E2" in ids({"age": 30, "sex": "F"}, ["xray"])


def test_all_exam_codes_known():
    for r in rulebook()["rules"]:
        def walk(c):
            if "exam" in c:
                for e in (c["exam"] if isinstance(c["exam"], list) else [c["exam"]]):
                    assert e in EXAMS, e
            for k in ("all", "any"):
                for s in c.get(k, []):
                    walk(s)
        walk(r["when"])


def test_stations_known():
    for r in rulebook()["rules"]:
        for s in r["stations"]:
            assert s in STATIONS, f"{r['id']}: 알 수 없는 검사실 {s}"


# ---------- 6. 일반 국가건강검진 규칙 ----------
@pytest.mark.parametrize("hours,expected", [(7.5, True), (8, False), (12, False)])
def test_fasting_threshold(hours, expected):
    assert ("Q1" in ids({"fasting_hours": hours}, CORE_EXAMS)) == expected


def test_fasting_missing_prompts():
    assert "Q0" in ids({"age": 40}, CORE_EXAMS)
    assert "Q0" not in ids({"age": 40}, ["xray"])  # 채혈이 없으면 묻지 않음


def test_menstruation_only_with_urine():
    assert "Q2" in ids({"sex": "F", "menstruating": True}, CORE_EXAMS)
    assert "Q2" not in ids({"sex": "F", "menstruating": True}, ["blood_draw"])


@pytest.mark.parametrize("sbp,dbp,expected", [(179, 109, False), (180, 90, True), (150, 110, True)])
def test_severe_hypertension(sbp, dbp, expected):
    assert ("P1" in ids({"sbp": sbp, "dbp": dbp}, CORE_EXAMS)) == expected


def test_extended_rules_silent_in_general_checkup():
    """일반검진 검사만 있으면 확장 규칙(조영제·MRI·내시경·운동부하)은 뜨지 않아야 한다."""
    ext = {r["id"] for r in rulebook()["rules"] if r.get("scope") == "extended"}
    f = {"implant": True, "egfr": 20, "osa": True, "cardiac_disease": True, "exertional_symptoms": True,
         "contrast_reaction": "moderate_severe", "claustrophobia": True, "sbp": 205, "dbp": 115,
         "height_cm": 160, "weight_kg": 100, "fasting_hours": 12}
    assert not ids(f, CORE_EXAMS) & ext


# ---------- 7. 검사 순서 추천 ----------
def _order(fields):
    return recommend_order(CORE_EXAMS, assess(fields, CORE_EXAMS))


def test_default_order():
    order, reasons = _order({"fasting_hours": 12})
    assert order == ["vitals", "vision_hearing", "urine", "blood_draw", "xray", "dental"] and not reasons


def test_hypoglycemia_blood_draw_first():
    order, _ = _order({"diabetes": True, "insulin": True, "fasting_hours": 12})
    assert order[:2] == ["vitals", "blood_draw"]


def test_syncope_seated_after_draw_xray_last():
    order, _ = _order({"prior_syncope_blood": True, "fasting_hours": 12})
    i = order.index("blood_draw")
    assert set(order[i + 1:i + 3]) == {"vision_hearing", "dental"} and order[-1] == "xray"
    assert order.index("urine") < i


def test_pregnancy_xray_last():
    order, _ = _order({"sex": "F", "age": 30, "pregnancy": "unknown", "fasting_hours": 12})
    assert order[-1] == "xray"


def test_order_keeps_all_exams():
    for f in [{"diabetes": True, "insulin": True, "prior_syncope_blood": True, "sex": "F", "age": 30}, {}]:
        order, _ = _order(f)
        assert sorted(order) == sorted(CORE_EXAMS)


# ---------- 8. 증례 검증 도구 ----------
def test_vignette_program_table():
    import vignette_validation as vv
    t = vv.program_table()
    assert len(t) == 30
    assert t.loc["V01", "실신 위험"] and t.loc["V02", "저혈당 위험"] and not t.loc["V08", vv.CATS].any()


def test_vignette_scoring(tmp_path):
    """평가지를 채운 상황을 흉내 내어 점수 계산이 동작하는지 확인."""
    import openpyxl
    import vignette_validation as vv
    path = tmp_path / "sheet.xlsx"
    vv.make_template(path)
    prog = vv.program_table()
    wb = openpyxl.load_workbook(path)
    for name in vv.SHEETS:
        ws = wb[name]
        for row in range(3, 33):
            cid = ws.cell(row=row, column=1).value
            for j, cat in enumerate(vv.CATS, 3):
                ws.cell(row=row, column=j).value = "O" if prog.at[cid, cat] else None
    wb["평가자B"].cell(row=3, column=3).value = None       # V01 실신 위험을 B가 놓쳤다고 가정
    wb["합의정답"].cell(row=10, column=3, value="O")       # V08에 정답 실신 위험이 있다고 가정 → 프로그램 FN
    wb.save(path)
    r = vv.score(path)
    assert 0 < r["rater_kappa"] < 1
    total = r["vs_truth"].set_index("범주").loc["전체"]
    assert total["FN"] == 1 and total["FP"] == 0
    assert len(r["disagreements"]) == 1
